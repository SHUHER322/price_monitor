"""Тесты выгрузки в Excel: собираем файл и читаем его обратно через openpyxl."""
import io
from datetime import datetime
from zoneinfo import ZoneInfo

from openpyxl import load_workbook

from app.db import ItemRow, Tracking
from app.excel import build_workbook, discount, export_filename, safe_filename, to_local

TZ = ZoneInfo("Asia/Tomsk")
NOW = datetime(2026, 10, 20, 19, 40, tzinfo=TZ)


def tracking(kind="search", title="стеклянный коврик") -> Tracking:
    return Tracking(id=1, user_id=111, kind=kind, source="wb", target=title, title=title,
                    target_price=None, created_at="2026-10-20 10:00:00")


def item(position, article, seller, price, old_price=650300) -> ItemRow:
    return ItemRow(position=position, article=article, name=f"Коврик {article}", brand="",
                   seller=seller, url=f"https://www.wildberries.ru/catalog/{article}/detail.aspx",
                   price=price, old_price=old_price, quantity=None, rating=4.9, feedbacks=9,
                   checked_at="2026-10-20 12:00:00")


ROWS = [item(1, "334082543", "GBS Техника", 361100),
        item(2, "594826793", "KA KA", 276300),
        item(3, "600232826", "KA KA", 276300),
        item(4, "1804906327", "RA", None)]          # нет в наличии


def open_book(rows=ROWS, kind="search"):
    return load_workbook(io.BytesIO(build_workbook(tracking(kind), rows, TZ, NOW)))


def test_items_sheet():
    sheet = open_book()["Товары"]
    assert sheet["A1"].value == "Место"
    assert sheet["F1"].value == "Цена"
    assert sheet["A2"].value == 1
    assert sheet["B2"].value == 334082543         # артикул — числом, удобно сортировать
    assert sheet["F2"].value == 3611              # цена в рублях
    assert sheet["H3"].value == 0.58              # скидка 58%
    assert sheet["L2"].hyperlink.target.endswith("334082543/detail.aspx")
    assert sheet["F5"].value is None              # нет в наличии — пустая цена
    assert sheet["H5"].value is None              # …и скидка пустая, без «—»
    assert sheet.auto_filter.ref == "A1:L5"
    assert sheet.freeze_panes == "A2"


def test_summary_sheet():
    book = open_book()
    assert book.sheetnames == ["Сводка", "Товары"]
    sheet = book["Сводка"]
    facts = {sheet.cell(row=r, column=1).value: sheet.cell(row=r, column=2).value
             for r in range(3, 11)}
    assert facts["Товаров"] == 4
    assert facts["В наличии"] == 3
    assert facts["Минимальная цена"] == 2763
    assert facts["Максимальная цена"] == 3611
    assert facts["Медиана цены"] == 2763
    assert facts["Цены на момент проверки"] == "20.10.2026 19:00"   # 12:00 UTC = 19:00 Томск


def test_sellers_table_for_search():
    sheet = open_book()["Сводка"]
    cells = [[c.value for c in row] for row in sheet.iter_rows(min_row=13, max_col=4)]
    header = cells.index(["Продавец", "Товаров в топе", "Лучшее место", "Средняя цена"])
    first = cells[header + 1]
    assert first[:3] == ["KA KA", 2, 2]           # больше всех товаров в топе


def test_no_sellers_table_for_seller():
    sheet = open_book(kind="seller")["Сводка"]
    values = [c.value for row in sheet.iter_rows() for c in row]
    assert "Продавцы в выдаче" not in values


def test_helpers():
    assert discount(item(1, "1", "x", 276300)) == 0.58
    assert discount(item(1, "1", "x", 276300, old_price=None)) is None
    assert safe_filename('коврик: "лучший" / 2026') == "коврик_лучший_2026"
    assert to_local(None, TZ) == ""
    assert export_filename(tracking(), NOW) == "поиск_wb_стеклянный_коврик_2026-10-20.xlsx"


def test_all_prices_in_summary_are_rubles():
    sheet = open_book()["Сводка"]
    formats = {sheet.cell(row=r, column=1).value: sheet.cell(row=r, column=2).number_format
               for r in range(3, 11)}
    for label in ("Минимальная цена", "Средняя цена", "Медиана цены", "Максимальная цена"):
        assert "₽" in formats[label], label


def test_stock_note_in_summary():
    sheet = open_book()["Сводка"]
    values = [c.value for row in sheet.iter_rows() for c in row if c.value]
    assert any("ограничивать" in str(v) for v in values)


def ozon_book(kind):
    t = Tracking(id=2, user_id=111, kind=kind, source="ozon", target="ridberg",
                 title="Ридберг", target_price=None, created_at="x")
    rows = [item(1, "1291636512", "Ридберг", 366800),
            ItemRow(2, "1291633069", "Микшер", "", "Ридберг", "https://ozon/1", 302300, None,
                    -1, None, 0, "2026-10-20 12:00:00")]
    return load_workbook(io.BytesIO(build_workbook(t, rows, TZ, NOW)))


def headers(sheet):
    return [c.value for c in sheet[1]]


def test_empty_values_are_dashes():
    sheet = open_book([item(1, "1", "x", 276300, old_price=None)])["Товары"]
    row = dict(zip(headers(sheet), [c.value for c in sheet[2]]))
    assert row["Цена до скидки"] == "—" and row["Скидка"] == "—"   # скидки нет


def test_ozon_seller_columns():
    book = ozon_book("seller")
    items = book["Товары"]
    titles = headers(items)
    assert "Бренд" not in titles                  # в каталоге Ozon бренда нет
    assert "Продавец" in titles
    assert "Цена по карте Ozon" in titles
    row = dict(zip(titles, [c.value for c in items[3]]))
    assert row["Остаток, шт."] == "в наличии"     # количество неизвестно
    assert row["Рейтинг"] == "—"                  # нет отзывов
    assert row["Скидка"] == "—"
    assert items.auto_filter.ref == "A1:K3"

    summary = book["Сводка"]
    assert summary["A1"].value == "Магазин (Ozon): Ридберг"
    values = [c.value for row in summary.iter_rows() for c in row if c.value]
    assert any("по карте Ozon" in str(v) for v in values)


def test_ozon_search_has_no_seller():
    book = ozon_book("search")
    assert "Продавец" not in headers(book["Товары"])
    values = [c.value for row in book["Сводка"].iter_rows() for c in row if c.value]
    assert "Продавцы в выдаче" not in values      # продавцов в выдаче Ozon нет

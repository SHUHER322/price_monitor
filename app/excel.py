"""Выгрузка отслеживания (магазин или поиск) в Excel.

Два листа:
    «Товары»  — все товары: место, цена, скидка, рейтинг, ссылка; с фильтрами
    «Сводка»  — итоги: сколько товаров, цены (мин/средняя/медиана/макс),
                а для поиска — какие продавцы занимают выдачу

Функция build_workbook чистая: на входе данные, на выходе байты файла.
Её легко тестировать, а отправкой файла в Telegram занимается обработчик.
"""
import io
import re
import statistics
from collections import defaultdict
from datetime import datetime
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .db import ItemRow, Tracking

RUB = '# ##0 "₽"'   # формат ячейки: 2 763 ₽
PERCENT = "0%"
HEADER_FONT = Font(bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor="2F3B52")
LINK_FONT = Font(color="0563C1", underline="single")
TITLE_FONT = Font(bold=True, size=14)

# Столбцы листа «Товары»: (заголовок, ширина)
COLUMNS = [
    ("Место", 7), ("Артикул", 12), ("Название", 50), ("Бренд", 18), ("Продавец", 22),
    ("Цена", 11), ("Цена до скидки", 14), ("Скидка", 11), ("Рейтинг", 11),
    ("Отзывы", 10), ("Остаток в регионе, шт.", 14), ("Ссылка", 12),
]


def to_local(utc_text: str | None, tz: ZoneInfo) -> str:
    """'2026-10-20 12:00:00' (UTC) → '20.10.2026 19:00' в поясе пользователя."""
    if not utc_text:
        return ""
    moment = datetime.strptime(utc_text, "%Y-%m-%d %H:%M:%S").replace(tzinfo=ZoneInfo("UTC"))
    return moment.astimezone(tz).strftime("%d.%m.%Y %H:%M")


def rubles(kopecks: int | None) -> float | None:
    return kopecks / 100 if kopecks is not None else None


def discount(row: ItemRow) -> float | None:
    """Скидка долей: 0.58 → в Excel покажется как 58%."""
    if not row.price or not row.old_price or row.old_price <= row.price:
        return None
    return round(1 - row.price / row.old_price, 2)


def safe_filename(text: str) -> str:
    """Убрать символы, которые нельзя в имени файла Windows."""
    cleaned = re.sub(r'[\\/:*?"<>|]+', "", text).strip()
    return re.sub(r"\s+", "_", cleaned)[:40] or "export"


def _header(sheet, row: int, titles: list[str]) -> None:
    for col, title in enumerate(titles, 1):
        cell = sheet.cell(row=row, column=col, value=title)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _items_sheet(sheet, rows: list[ItemRow]) -> None:
    sheet.title = "Товары"
    _header(sheet, 1, [title for title, _ in COLUMNS])
    for col, (_, width) in enumerate(COLUMNS, 1):
        sheet.column_dimensions[get_column_letter(col)].width = width

    for r, item in enumerate(rows, 2):
        values = [item.position, int(item.article) if item.article.isdigit() else item.article, item.name, item.brand, item.seller,
                  rubles(item.price), rubles(item.old_price), discount(item),
                  item.rating, item.feedbacks, item.quantity, "открыть"]
        for col, value in enumerate(values, 1):
            sheet.cell(row=r, column=col, value=value)
        sheet.cell(row=r, column=6).number_format = RUB
        sheet.cell(row=r, column=7).number_format = RUB
        sheet.cell(row=r, column=8).number_format = PERCENT
        link = sheet.cell(row=r, column=12)
        link.hyperlink = item.url
        link.font = LINK_FONT

    sheet.freeze_panes = "A2"   # заголовок не уезжает при прокрутке
    if rows:
        # Фильтры в заголовке: сортировать по цене, выбрать продавца и т.д.
        sheet.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}{len(rows) + 1}"


def _summary_sheet(sheet, tracking: Tracking, rows: list[ItemRow], tz: ZoneInfo,
                   exported_at: str) -> None:
    kind = {"seller": "Магазин", "search": "Поиск"}.get(tracking.kind, tracking.kind)
    sheet["A1"] = f"{kind}: {tracking.title}"
    sheet["A1"].font = TITLE_FONT
    sheet.column_dimensions["A"].width = 32
    sheet.column_dimensions["B"].width = 16
    sheet.column_dimensions["C"].width = 16
    sheet.column_dimensions["D"].width = 14

    prices = [r.price / 100 for r in rows if r.price is not None]
    checked = max((r.checked_at for r in rows if r.checked_at), default=None)
    # (подпись, значение, это деньги?) — формат «₽» задаём явно, а не по словам
    # в подписи: так «Медиана цены» однажды осталась без рублей
    facts = [
        ("Выгружено", exported_at, False),
        ("Цены на момент проверки", to_local(checked, tz), False),
        ("Товаров", len(rows), False),
        ("В наличии", sum(1 for r in rows if r.price is not None), False),
        ("Минимальная цена", min(prices) if prices else None, True),
        ("Средняя цена", round(statistics.mean(prices)) if prices else None, True),
        # Медиана — «серединная» цена: половина товаров дешевле, половина дороже.
        # Не так сильно, как среднее, зависит от одного очень дорогого товара
        ("Медиана цены", round(statistics.median(prices)) if prices else None, True),
        ("Максимальная цена", max(prices) if prices else None, True),
    ]
    for i, (label, value, money) in enumerate(facts, 3):
        sheet.cell(row=i, column=1, value=label).font = Font(bold=True)
        cell = sheet.cell(row=i, column=2, value=value)
        if money:
            cell.number_format = RUB
    note = sheet.cell(row=len(facts) + 3, column=1, value=(
        "Остаток — со склада, ближайшего к региону доставки. WB может ограничивать "
        "показываемое число: одинаковое максимальное значение у многих товаров "
        "значит «столько и больше»."))
    note.font = Font(italic=True, color="666666")

    if tracking.kind != "search" or not rows:
        return
    # Для поиска: кто из продавцов занимает выдачу
    start = len(facts) + 6
    sheet.cell(row=start - 1, column=1, value="Продавцы в выдаче").font = Font(bold=True, size=12)
    _header(sheet, start, ["Продавец", "Товаров в топе", "Лучшее место", "Средняя цена"])
    by_seller: dict[str, list[ItemRow]] = defaultdict(list)
    for r in rows:
        by_seller[r.seller or "—"].append(r)
    ranking = sorted(by_seller.items(), key=lambda kv: (-len(kv[1]),
                     min(x.position or 10**6 for x in kv[1])))
    for i, (seller, items) in enumerate(ranking, start + 1):
        seller_prices = [x.price / 100 for x in items if x.price is not None]
        sheet.cell(row=i, column=1, value=seller)
        sheet.cell(row=i, column=2, value=len(items))
        sheet.cell(row=i, column=3, value=min((x.position for x in items if x.position),
                                              default=None))
        avg = sheet.cell(row=i, column=4,
                         value=round(statistics.mean(seller_prices)) if seller_prices else None)
        avg.number_format = RUB


def build_workbook(tracking: Tracking, rows: list[ItemRow], tz: ZoneInfo,
                   now: datetime | None = None) -> bytes:
    """Собрать файл Excel и вернуть его содержимое (байты)."""
    now = now or datetime.now(tz)
    workbook = Workbook()
    _items_sheet(workbook.active, rows)
    _summary_sheet(workbook.create_sheet("Сводка", 0), tracking, rows, tz,
                   now.strftime("%d.%m.%Y %H:%M"))
    workbook.active = 0   # открываться файл будет на сводке
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def export_filename(tracking: Tracking, now: datetime) -> str:
    prefix = {"seller": "магазин", "search": "поиск"}.get(tracking.kind, tracking.kind)
    return f"{prefix}_{safe_filename(tracking.title)}_{now:%Y-%m-%d}.xlsx"
